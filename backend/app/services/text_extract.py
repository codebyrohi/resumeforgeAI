import io


def extract_text(kind: str, data: bytes) -> str:
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if kind == "docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines).strip()
    raise ValueError(f"unsupported kind: {kind}")
